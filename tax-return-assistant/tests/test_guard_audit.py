import json
import logging

import openpyxl
import pytest

from tra.privacy import guard
from tra.privacy.audit import AuditLog
from tra.privacy.guard import FileLockedError, SensitiveDataBlocked, atomic_write_text, check_rows, check_text, check_workbook
from tra.privacy.logging_filter import configure_logging


@pytest.fixture
def samples(make_tfn, make_medicare):
    return {
        "TFN": make_tfn(),
        "MEDICARE": make_medicare(),
        "CARD": "4111 1111 1111 1111",
        "BSB": "BSB 062-000",
        "ACCOUNT": "Account number 1234 5678",
        "LICENCE": "Driver licence no. 12345678",
        "PASSPORT": "Passport PA1234567",
    }


def test_check_text_blocks_each_type(samples):
    for kind, sample in samples.items():
        with pytest.raises(SensitiveDataBlocked) as err:
            check_text(f"note: {sample}", "figures.csv")
        assert kind in str(err.value)
        assert sample not in str(err.value)
        assert "figures.csv" in str(err.value)


def test_check_text_allows_normal_output():
    check_text("SMIT01_FY2026_DIVSTMT_2026-03-24.pdf", "filename")
    check_text("Interest 412.35 ingested 2026-10-14T05:49:00+10:00", "index.csv")


def test_check_rows_blocks_int_tfn(make_tfn):
    with pytest.raises(SensitiveDataBlocked):
        check_rows([["F0001", "income.interest", int(make_tfn())]], "figures.csv")
    check_rows([["F0001", "income.interest", 412.35, 85000]], "figures.csv")


def test_check_workbook_blocks_cell(make_tfn):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Interest"
    ws["A1"] = "=SUM(B1:B5)"
    ws["B7"] = f"Account holder TFN {make_tfn()}"
    with pytest.raises(SensitiveDataBlocked) as err:
        check_workbook(wb, "workpaper")
    assert "Interest!B7" in str(err.value)


def test_atomic_write_writes(tmp_path):
    target = tmp_path / "figures.csv"
    atomic_write_text(target, "a,b\n", "figures.csv")
    assert target.read_text() == "a,b\n"


def test_atomic_write_permission_error(tmp_path, monkeypatch):
    target = tmp_path / "figures.csv"
    target.write_text("original\n")

    def locked(*_args):
        raise PermissionError("locked")

    monkeypatch.setattr(guard.os, "replace", locked)
    with pytest.raises(FileLockedError) as err:
        atomic_write_text(target, "new\n", "figures.csv")
    assert "figures.csv is open in another program" in str(err.value)
    assert target.read_text() == "original\n"
    assert list(tmp_path.iterdir()) == [target]


def test_atomic_write_refuses_sensitive(tmp_path, make_tfn):
    target = tmp_path / "state.json"
    with pytest.raises(SensitiveDataBlocked):
        atomic_write_text(target, json.dumps({"note": make_tfn()}), "state.json")
    assert not target.exists()


def test_audit_records_and_blocks(tmp_path, make_tfn):
    log = AuditLog(tmp_path / "audit" / "audit-2026.jsonl")
    log.record("document_ingested", client="SMIT01", doc_id="D001", redaction_counts={"TFN": 1})
    tfn = make_tfn()
    log.record("document_ingested", client="SMIT01", note=f"tfn {tfn}")
    lines = [json.loads(line) for line in log.path.read_text().splitlines()]
    assert lines[0]["action"] == "document_ingested"
    assert lines[0]["client"] == "SMIT01"
    assert lines[0]["user"]
    assert "T" in lines[0]["ts"]
    assert lines[1] == {**lines[1], "action": "guard_blocked", "blocked_action": "document_ingested", "detector": "TFN"}
    assert tfn not in log.path.read_text()


def test_logging_filter_redacts(tmp_path, make_tfn):
    log_path = tmp_path / "logs" / "tra.log"
    configure_logging(log_path)
    tfn = make_tfn()
    logging.getLogger("tra.test").warning("client TFN %s", tfn)
    for handler in logging.getLogger("tra").handlers:
        handler.flush()
    content = log_path.read_text()
    assert "[TFN REDACTED]" in content
    assert tfn not in content
