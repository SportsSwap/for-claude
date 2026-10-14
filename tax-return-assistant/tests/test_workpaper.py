import shutil
import subprocess
from decimal import Decimal

import openpyxl
import pytest
from wpctx import demo_context

from tra.privacy.guard import FileLockedError, SensitiveDataBlocked
from tra.workpaper.build import build_workpaper


@pytest.fixture
def built(tmp_path):
    ctx = demo_context(tmp_path)
    path = tmp_path / "wp.xlsx"
    wp_map = build_workpaper(path, ctx)
    return path, wp_map, ctx


def test_tabs_present(built):
    path, _, _ = built
    names = openpyxl.load_workbook(path).sheetnames
    for tab in ("Cover", "Summary", "Salary & Wages", "Interest", "Dividends", "Rental", "D5 Other work-related",
                "Private Health", "Tax Calc", "Rates", "Sources", "Review Notes"):
        assert tab in names
    assert names[0] == "Cover" and names[-1] == "Review Notes"


def test_totals_are_formulas(built):
    path, wp_map, _ = built
    wb = openpyxl.load_workbook(path)
    refs = wp_map.total_cells + list(wp_map.summary_cells.values()) + list(wp_map.taxcalc_cells.values())
    assert len(wp_map.taxcalc_cells) == 17
    assert "income.interest" in wp_map.summary_cells
    for ref in refs:
        sheet, cell = ref.rsplit("!", 1)
        value = wb[sheet.strip("'")][cell].value
        assert isinstance(value, str) and value.startswith("="), ref


def test_every_figure_row_has_source(built):
    path, _, ctx = built
    wb = openpyxl.load_workbook(path)
    rows = 0
    for tab in ("Salary & Wages", "Interest", "Dividends", "Rental", "D5 Other work-related", "Private Health"):
        ws = wb[tab]
        for row in ws.iter_rows(min_row=5):
            if row[0].value and str(row[0].value).startswith("F"):
                rows += 1
                assert row[4].value, (tab, row[0].value)
                assert row[5].value, (tab, row[0].value)
                assert row[4].hyperlink is not None
    assert rows == len(ctx.figures)


def test_banner_and_rates_status(built):
    path, _, _ = built
    wb = openpyxl.load_workbook(path)
    assert wb["Cover"]["A1"].value == "DRAFT – NOT FOR LODGMENT"
    assert "A1:B2" in [str(r) for r in wb["Cover"].merged_cells.ranges]
    assert all(ws.page_setup.fitToWidth == 1 for ws in wb.worksheets)
    texts = [c.value for row in wb["Rates"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("unverified" in t for t in texts)


def test_review_notes_ranked_with_citations(built):
    path, _, ctx = built
    ws = openpyxl.load_workbook(path)["Review Notes"]
    risks = [ws.cell(row=r, column=2).value for r in range(5, 5 + len(ctx.notes))]
    assert risks == [n.risk for n in ctx.notes]
    citations = [ws.cell(row=r, column=8).value for r in range(5, 5 + len(ctx.notes))]
    assert any(c and "s900-35" in c and "to be confirmed" in c for c in citations)


def test_guard_blocks_leak(tmp_path, make_tfn):
    tfn = make_tfn()
    ctx = demo_context(tmp_path)
    ctx.figures[0].description = f"note {tfn}"
    path = tmp_path / "leak.xlsx"
    with pytest.raises(SensitiveDataBlocked):
        build_workpaper(path, ctx)
    assert not path.exists()
    assert not list(tmp_path.glob("leak.xlsx*"))


def test_locked_file(tmp_path, monkeypatch):
    ctx = demo_context(tmp_path)

    def locked(*_args):
        raise PermissionError("locked")

    monkeypatch.setattr("tra.workpaper.build.os.replace", locked)
    with pytest.raises(FileLockedError):
        build_workpaper(tmp_path / "wp.xlsx", ctx)


def _recalculate(path, out_dir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(out_dir / path.name, data_only=True)


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_recalc_matches_calc_a(built, tmp_path):
    path, wp_map, ctx = built
    wb = _recalculate(path, tmp_path / "recalc")
    expected = ctx.calc_a.components()
    for name, ref in wp_map.taxcalc_cells.items():
        sheet, cell = ref.rsplit("!", 1)
        value = wb[sheet.strip("'")][cell].value
        assert value is not None, name
        assert abs(Decimal(str(value)) - expected[name]) <= Decimal("0.01"), (name, value, expected[name])


def _set_salary(amount):
    def mutate(figures):
        for f in figures:
            if f.label_key == "income.salary" and f.confirmed:
                f.amount = Decimal(amount)
    return mutate


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
@pytest.mark.parametrize("case", ["help_mls_cgt", "senior_low_income"])
def test_recalc_matches_calc_a_other_branches(tmp_path, case):
    from tra.models import ClientState
    if case == "help_mls_cgt":
        def mutate(figures):
            _set_salary("140000")(figures)
            from tra.models import Figure
            figures.append(Figure("F0900", "ded.d12", "Personal super contribution", Decimal("10000"), None, "D002",
                                  1, "manual", True, True, ""))

        ctx = demo_context(tmp_path, state=ClientState("DEMO01", 2026, help_debt=True, hospital_cover_days=200),
                           mutate=mutate,
                           cgt_rows=[["C001", "BHP shares", "2020-01-10", "2026-03-24", "15000.00", "9000.00", "D004", "1", "Y"],
                                     ["C002", "Crypto", "2025-10-01", "2026-02-01", "3000.00", "1000.00", "", "", "Y"]])
        assert ctx.calc_a.help > 0 and ctx.calc_a.mls > 0 and ctx.draft.amount("cgt.net_gain") > 0
        assert ctx.draft.amount("ded.d12") == Decimal("10000")
    else:
        ctx = demo_context(tmp_path, state=ClientState("DEMO01", 2026, sapto_status="single", age_band="70plus"),
                           mutate=_set_salary("12000"))
        assert ctx.calc_a.sapto > 0 and ctx.calc_a.lito > 0
    path = tmp_path / "wp.xlsx"
    wp_map = build_workpaper(path, ctx)
    wb = _recalculate(path, tmp_path / "recalc")
    expected = ctx.calc_a.components()
    for name, ref in wp_map.taxcalc_cells.items():
        sheet, cell = ref.rsplit("!", 1)
        value = wb[sheet.strip("'")][cell].value
        assert value is not None, name
        assert abs(Decimal(str(value)) - expected[name]) <= Decimal("0.01"), (case, name, value, expected[name])


def test_no_formula_sums_over_its_own_cell(tmp_path):
    import re
    from openpyxl.utils import range_boundaries
    ctx = demo_context(tmp_path, cgt_rows=[["C001", "BHP shares", "2020-01-10", "2026-03-24", "15000.00", "9000.00",
                                            "D004", "1", "Y"]])
    path = tmp_path / "wp.xlsx"
    build_workpaper(path, ctx)
    for ws in openpyxl.load_workbook(path).worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=SUM("):
                    for rng in re.findall(r"SUM\(([A-Z]+\d+:[A-Z]+\d+)\)", cell.value):
                        min_col, min_row, max_col, max_row = range_boundaries(rng)
                        lo, hi = sorted((min_row, max_row))
                        assert not (min_col <= cell.column <= max_col and lo <= cell.row <= hi), \
                            f"{ws.title}!{cell.coordinate} {cell.value}"


def test_no_total_mixes_different_labels(built):
    import re
    path, _, _ = built
    wb = openpyxl.load_workbook(path)
    for name in ("Salary & Wages", "Interest", "Dividends", "Rental", "D5 Other work-related", "Private Health"):
        ws = wb[name]
        sums = 0
        for row in ws.iter_rows():
            for cell in row:
                m = re.fullmatch(r"=SUM\(J(\d+):J(\d+)\)", str(cell.value or ""))
                if m:
                    sums += 1
                    keys = {ws.cell(row=r, column=11).value for r in range(int(m.group(1)), int(m.group(2)) + 1)} - {None}
                    assert len(keys) == 1, f"{name}!{cell.coordinate} adds {keys}"
        assert sums >= 1, name


def test_rental_tab_shows_net_rent(built):
    path, _, ctx = built
    ws = openpyxl.load_workbook(path)["Rental"]
    labels = [ws.cell(row=r, column=3).value for r in range(1, ws.max_row + 1)]
    assert "Net rent" in labels
