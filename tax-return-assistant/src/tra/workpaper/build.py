"""Build the Excel workpaper (spec §9).

Figures and rates are the only typed-in numbers; every total, the lead
schedule and the whole tax calculation are live formulas, so the reviewer can
follow the working in Excel. Calculation B's results sit beside the formulas
as a check. Layout and styling come from layout.py.
"""

from __future__ import annotations

import os
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..checks import label_name
from ..dates import days_in_fy, fy_bounds, fy_label
from ..labels import LabelDef
from ..models import ClientState, Document, Figure, ReviewNote
from ..privacy.guard import FileLockedError, check_workbook
from ..return_.cgt import discount_eligible
from ..return_.prepare import DraftReturn
from ..taxcalc.calc_a import DEDUCTION_KEYS
from ..taxcalc.rates import Rates
from ..taxcalc.result import TaxResult
from . import layout as L

BANNER = "DRAFT – NOT FOR LODGMENT"
CGT_TAB = "Capital Gains"
FIRST_ROW = 5  # data starts below title, subtitle, blank and header rows

TAXCALC_ROWS = [
    ("total_income", "Total income"),
    ("deductions", "Total deductions"),
    ("taxable_income", "Taxable income (cents dropped)"),
    ("basic_tax", "Tax on taxable income"),
    ("lito", "Low income tax offset"),
    ("sapto", "Seniors and pensioners tax offset"),
    ("offsets_used", "Non-refundable offsets used (limited to tax)"),
    ("medicare_levy", "Medicare levy"),
    ("mls_income", "Income for MLS purposes"),
    ("mls", "Medicare levy surcharge"),
    ("repayment_income", "Repayment income (HELP)"),
    ("help", "HELP compulsory repayment"),
    ("phi_adjustment", "Private health insurance rebate adjustment"),
    ("franking_credits", "Less: franking credits (refundable)"),
    ("payg_withheld", "Less: PAYG withheld"),
    ("tfn_withheld", "Less: TFN amounts withheld"),
    ("net_result", "Estimated payable / (refund)"),
]


@dataclass
class WorkpaperContext:
    client_code: str
    income_year: int
    state: ClientState
    documents: list[Document]
    figures: list[Figure]  # confirmed only
    draft: DraftReturn
    rates: Rates
    calc_a: TaxResult
    calc_b: TaxResult
    notes: list[ReviewNote]
    prefill: dict[str, Decimal]
    prior_year: dict[str, Decimal]
    labels: dict[str, LabelDef]
    prepared_by: str
    tool_version: str
    cgt_events: list = field(default_factory=list)
    citations: dict = field(default_factory=dict)


@dataclass
class WorkpaperMap:
    total_cells: list[str] = field(default_factory=list)
    summary_cells: dict[str, str] = field(default_factory=dict)
    taxcalc_cells: dict[str, str] = field(default_factory=dict)


def _q(sheet: str) -> str:
    return f"'{sheet}'"


def _link(location: str) -> str:
    if location.startswith(("http://", "https://")):
        return location
    path = Path(location)
    return path.as_uri() if path.is_absolute() else location


def _title(ws: Worksheet, title: str, subtitle: str) -> None:
    ws["A1"] = title
    L.style_title(ws["A1"])
    ws["A2"] = subtitle
    L.style_subtitle(ws["A2"])


# --- category tabs --------------------------------------------------------------------------

FIGURE_HEADERS = ["Figure ID", "Date", "Description", "Return label", "Source document", "Page", "Method",
                  "Confirmed", "Substantiated", "Amount", "key", "Notes"]


def _write_header(ws: Worksheet, row: int) -> None:
    for col, header in enumerate(FIGURE_HEADERS, start=1):
        ws.cell(row=row, column=col, value=header)
    L.style_header_row(ws, row, len(FIGURE_HEADERS))


def _write_rows(ws: Worksheet, figures: list[Figure], ctx: WorkpaperContext, docs: dict[str, Document],
                row: int) -> int:
    for f in figures:
        d = ctx.labels[f.label_key]
        ws.cell(row=row, column=1, value=f.figure_id)
        if f.date:
            ws.cell(row=row, column=2, value=f.date).number_format = L.DATE
        ws.cell(row=row, column=3, value=f.description)
        ws.cell(row=row, column=4, value=f"{d.item} {d.label}".strip() if len(d.label) <= 2 else d.item)
        doc = docs.get(f.doc_id or "")
        source = ws.cell(row=row, column=5)
        if doc:
            source.value = doc.filed_name
            source.hyperlink = _link(doc.filed_location)
            L.style_link(source)
        else:
            source.value = "NO SOURCE" if not f.doc_id else f"{f.doc_id} (not in index)"
            L.style_warning(source)
        ws.cell(row=row, column=6, value=f.page if f.page is not None else ("" if doc else "–"))
        ws.cell(row=row, column=7, value=f.method)
        ws.cell(row=row, column=8, value="Y" if f.confirmed else "N")
        ws.cell(row=row, column=9, value="" if f.substantiated is None else ("Y" if f.substantiated else "N"))
        L.style_input(ws.cell(row=row, column=10, value=f.amount))
        ws.cell(row=row, column=11, value=f.label_key)
        ws.cell(row=row, column=12, value=f.notes)
        row += 1
    return row


def _write_figure_rows(ws: Worksheet, figures: list[Figure], ctx: WorkpaperContext, docs: dict[str, Document],
                       start: int) -> int:
    _write_header(ws, start - 1)
    return _write_rows(ws, figures, ctx, docs, start)


RENTAL_DEDUCTION_KEYS = ("rent.interest", "rent.capital_works", "rent.other")


def _category_tab(ws: Worksheet, tab: str, figures: list[Figure], ctx: WorkpaperContext,
                  docs: dict[str, Document], wp: WorkpaperMap) -> None:
    """Figures grouped by return label, each group with its own subtotal (labels are never added together)."""
    _title(ws, f"{tab} — {ctx.client_code} {fy_label(ctx.income_year)}",
           "Confirmed figures only. Each amount traces to the filed source document and page.")
    _write_header(ws, FIRST_ROW - 1)
    row = FIRST_ROW
    subtotals: dict[str, str] = {}
    for key in [k for k in ctx.labels if any(f.label_key == k for f in figures)]:
        start = row
        row = _write_rows(ws, [f for f in figures if f.label_key == key], ctx, docs, row)
        ws.cell(row=row, column=3, value=f"Total {label_name(ctx.labels, key)}")
        L.style_label(ws.cell(row=row, column=3))
        L.style_total(ws.cell(row=row, column=10, value=f"=SUM(J{start}:J{row - 1})"))
        wp.total_cells.append(f"{_q(tab)}!J{row}")
        subtotals[key] = f"J{row}"
        row += 2
    if "rent.gross" in subtotals or any(k in subtotals for k in RENTAL_DEDUCTION_KEYS):
        less = "".join(f"-{subtotals[k]}" for k in RENTAL_DEDUCTION_KEYS if k in subtotals)
        ws.cell(row=row, column=3, value="Net rent")
        L.style_label(ws.cell(row=row, column=3))
        L.style_total(ws.cell(row=row, column=10, value=f"={subtotals.get('rent.gross', '0')}{less}"))
        wp.total_cells.append(f"{_q(tab)}!J{row}")
    ws.column_dimensions["K"].hidden = True
    ws.freeze_panes = f"A{FIRST_ROW}"
    L.set_widths(ws, L.CATEGORY_WIDTHS)
    L.setup_print(ws, header_row=FIRST_ROW - 1)


def _cgt_tab(ws: Worksheet, figures: list[Figure], ctx: WorkpaperContext, docs: dict[str, Document],
             wp: WorkpaperMap) -> dict[str, str]:
    """Prior-loss figures, the CGT event schedule and the s102-5 method statement, as formulas."""
    _title(ws, f"{CGT_TAB} — {ctx.client_code} {fy_label(ctx.income_year)}",
           "Net capital gain per the ITAA 1997 s102-5 method statement: losses reduce non-discount gains first.")
    end = _write_figure_rows(ws, figures, ctx, docs, FIRST_ROW)
    ws.cell(row=end, column=3, value="Capital losses brought forward")
    # No prior-loss figures: nothing to sum (a SUM over the empty range would include this cell).
    prior_losses = ws.cell(row=end, column=10, value=f"=SUM(J{FIRST_ROW}:J{end - 1})" if end > FIRST_ROW else "=0")
    L.style_total(prior_losses)
    wp.total_cells.append(f"{_q(CGT_TAB)}!J{end}")
    ws.column_dimensions["K"].hidden = True

    head = end + 3
    headers = ["Event", "Asset", "Acquired", "Disposed", "Proceeds", "Cost base", "Gain / (loss)",
               "Discount eligible", "Source document", "Page"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=head, column=col, value=h)
    L.style_header_row(ws, head, len(headers))
    row = head + 1
    events = [e for e in ctx.cgt_events if e.confirmed]
    for e in events:
        ws.cell(row=row, column=1, value=e.event_id)
        ws.cell(row=row, column=2, value=e.asset)
        ws.cell(row=row, column=3, value=e.acquired).number_format = L.DATE
        ws.cell(row=row, column=4, value=e.disposed).number_format = L.DATE
        L.style_input(ws.cell(row=row, column=5, value=e.proceeds))
        L.style_input(ws.cell(row=row, column=6, value=e.cost_base))
        L.style_formula(ws.cell(row=row, column=7, value=f"=E{row}-F{row}"))
        ws.cell(row=row, column=8, value="Y" if discount_eligible(e.acquired, e.disposed) else "N")
        doc = docs.get(e.doc_id or "")
        ws.cell(row=row, column=9, value=doc.filed_name if doc else "NO SOURCE")
        ws.cell(row=row, column=10, value=e.page)
        row += 1
    first, last = head + 1, max(head + 1, row - 1)
    g, h = f"G{first}:G{last}", f"H{first}:H{last}"
    half = "0.5" if ctx.state.resident else "1"
    calc = [
        ("total", "Total current year capital gains (18 H)", f'=SUMIF({g},">0")'),
        ("nd", "Gains not eligible for the discount", f'=SUMIFS({g},{g},">0",{h},"N")'),
        ("dg", "Gains eligible for the discount", f'=SUMIFS({g},{g},">0",{h},"Y")'),
        ("losses", "Capital losses: this year plus brought forward", f'=-SUMIF({g},"<0")+J{end}'),
        ("nd_after", "Non-discount gains after losses", "=MAX(0,{nd}-{losses})"),
        ("dg_after", "Discount gains after remaining losses", "=MAX(0,{dg}-MAX(0,{losses}-{nd}))"),
        ("net", f"Net capital gain (18 A){' — 50% discount applied' if ctx.state.resident else ''}",
         "=ROUND({nd_after}+{dg_after}*" + half + ",2)"),
    ]
    refs: dict[str, str] = {}
    row += 2
    for key, label, formula in calc:
        refs[key] = f"J{row}"
        ws.cell(row=row, column=3, value=label)
        cell = ws.cell(row=row, column=10, value=formula.format(**refs))
        (L.style_total if key in ("total", "net") else L.style_formula)(cell)
        wp.total_cells.append(f"{_q(CGT_TAB)}!J{row}")
        row += 1
    L.set_widths(ws, L.CATEGORY_WIDTHS)
    L.setup_print(ws)
    return {"cgt.total_gains": f"{_q(CGT_TAB)}!$J${refs['total'][1:]}",
            "cgt.net_gain": f"{_q(CGT_TAB)}!$J${refs['net'][1:]}"}


# --- summary (lead schedule) ----------------------------------------------------------------

def _summary(ws: Worksheet, ctx: WorkpaperContext, tab_of: dict[str, str], cgt_refs: dict[str, str],
             wp: WorkpaperMap) -> dict[str, str]:
    _title(ws, f"Lead schedule — {ctx.client_code} {fy_label(ctx.income_year)}",
           "Every return label in ATO item order. This year links to the working tabs; pre-fill and prior year are inputs.")
    headers = ["Item", "Label", "Description", "This year", "Pre-fill", "Difference", "Prior year", "Movement",
               "Review notes", "key"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=FIRST_ROW - 1, column=col, value=h)
    L.style_header_row(ws, FIRST_ROW - 1, len(headers))
    per_label = defaultdict(Counter)
    for n in ctx.notes:
        if n.label_key:
            per_label[n.label_key][n.risk] += 1
    refs: dict[str, str] = {}
    row = FIRST_ROW
    for key, d in ctx.labels.items():
        if d.role == "info" and not d.computed:
            continue
        ws.cell(row=row, column=1, value=d.item)
        ws.cell(row=row, column=2, value=d.label if len(d.label) <= 2 else "")
        ws.cell(row=row, column=3, value=d.description)
        this_year = ws.cell(row=row, column=4)
        if key in cgt_refs:
            this_year.value = f"={cgt_refs[key]}"
        elif key in tab_of:
            tab = _q(tab_of[key])
            this_year.value = f'=SUMIFS({tab}!$J:$J,{tab}!$K:$K,"{key}")'
        else:
            this_year.value = 0
        if isinstance(this_year.value, str):
            L.style_formula(this_year)
            wp.summary_cells[key] = f"Summary!D{row}"
        else:
            L.style_input(this_year)
        if key in ctx.prefill:
            L.style_input(ws.cell(row=row, column=5, value=ctx.prefill[key]))
        L.style_formula(ws.cell(row=row, column=6, value=f'=IF(E{row}="","",D{row}-E{row})'))
        if key in ctx.prior_year:
            L.style_input(ws.cell(row=row, column=7, value=ctx.prior_year[key]))
        L.style_formula(ws.cell(row=row, column=8, value=f'=IF(OR(G{row}="",G{row}=0),"",(D{row}-G{row})/G{row})'),
                        L.PERCENT)
        counts = per_label.get(key)
        if counts:
            ws.cell(row=row, column=9, value=", ".join(f"{counts[r]} {r}" for r in ("High", "Medium", "Low") if counts[r]))
        ws.cell(row=row, column=10, value=key)
        refs[key] = f"Summary!$D${row}"
        row += 1
    ws.column_dimensions["J"].hidden = True
    ws.freeze_panes = f"A{FIRST_ROW}"
    L.set_widths(ws, L.SUMMARY_WIDTHS)
    L.setup_print(ws, header_row=FIRST_ROW - 1)
    return refs


# --- rates ---------------------------------------------------------------------------------

class _Block:
    """Writes label/value rows down a sheet and remembers where each value went."""

    def __init__(self, ws: Worksheet, row: int):
        self.ws, self.row, self.refs = ws, row, {}

    def heading(self, text: str) -> None:
        self.row += 1
        self.ws.cell(row=self.row, column=1, value=text)
        L.style_label(self.ws.cell(row=self.row, column=1))
        self.row += 1

    def value(self, name: str, label: str, value, fmt: str = L.NUMBER) -> None:
        self.ws.cell(row=self.row, column=1, value=label)
        L.style_input(self.ws.cell(row=self.row, column=2, value=value), fmt)
        self.refs[name] = f"Rates!$B${self.row}"
        self.row += 1

    def table(self, name: str, headers: list[str], rows: list[list], fmts: list[str]) -> None:
        for col, h in enumerate(headers, start=1):
            self.ws.cell(row=self.row, column=col, value=h)
        L.style_header_row(self.ws, self.row, len(headers))
        first = self.row + 1
        for values in rows:
            self.row += 1
            for col, (v, fmt) in enumerate(zip(values, fmts), start=1):
                L.style_input(self.ws.cell(row=self.row, column=col, value=v), fmt)
        last = self.row
        for col in range(1, len(headers) + 1):
            letter = chr(ord("A") + col - 1)
            self.refs[f"{name}.{col}"] = f"Rates!${letter}${first}:${letter}${last}"
        self.refs[f"{name}.row1"] = f"Rates!$B${first}:${chr(ord('A') + len(headers) - 1)}${first}"
        self.row += 1


def _rates(ws: Worksheet, ctx: WorkpaperContext) -> dict[str, str]:
    r, st = ctx.rates, ctx.state
    _title(ws, f"Rates and thresholds — FY{r.income_year}",
           "Copied from config/rates so this workbook stands alone. Blue = input.")
    status = ws.cell(row=3, column=1, value=(
        f"Status: {r.status}" + (f" (verified by {r.verified_by} on {r.verified_on})" if r.status == "verified"
                                 else " — every figure must be checked against ato.gov.au by a registered agent")))
    if r.status != "verified":
        L.style_warning(status)
    b = _Block(ws, 4)
    b.heading("Resident tax rates")
    b.table("brackets", ["Income over", "Marginal rate", "Tax on income up to threshold"],
            [[x.from_, x.rate, x.base] for x in r.brackets], [L.WHOLE, L.PERCENT, L.NUMBER])
    b.heading("Low income tax offset")
    b.value("lito_max", "Maximum offset", r.lito.max)
    b.value("lito_full_to", "Full offset up to", r.lito.full_to, L.WHOLE)
    b.value("lito_t1_rate", "First taper rate", r.lito.taper1_rate, L.RATE)
    b.value("lito_t1_to", "First taper ends at", r.lito.taper1_to, L.WHOLE)
    b.value("lito_t2_base", "Offset at end of first taper", r.lito.taper2_base)
    b.value("lito_t2_rate", "Second taper rate", r.lito.taper2_rate, L.RATE)
    sapto = st.sapto_status != "none"
    threshold = r.medicare.sapto_single if sapto else r.medicare.single
    b.heading(f"Medicare levy ({'SAPTO single' if sapto else 'single'} thresholds)")
    b.value("ml_rate", "Levy rate", r.medicare.rate, L.PERCENT)
    b.value("ml_shade", "Shade-in rate over the lower threshold", r.medicare.shade_in, L.PERCENT)
    b.value("ml_lower", "Lower threshold", threshold.lower, L.WHOLE)
    b.value("ml_upper", "Full levy from", threshold.upper, L.WHOLE)
    if sapto:
        s = r.sapto[st.sapto_status]
        b.heading(f"Seniors and pensioners tax offset ({st.sapto_status.replace('_', ' ')})")
        b.value("sapto_max", "Maximum offset", s.max)
        b.value("sapto_shade", "Shade-out threshold", s.shade_out, L.WHOLE)
        b.value("sapto_cut", "Cut-out threshold", s.cut_out, L.WHOLE)
        b.value("sapto_taper", "Taper rate", r.sapto_taper, L.PERCENT)
    family = st.has_spouse or st.dependants > 0
    uplift = r.mls.family_per_child_after_first * max(0, st.dependants - 1) if family else Decimal("0")
    tiers = r.mls.family if family else r.mls.single
    b.heading(f"Medicare levy surcharge and PHI rebate tiers ({'family' if family else 'single'})")
    b.table("mls", ["Income over", "MLS rate"], [[t.from_ + uplift if t.from_ else t.from_, t.rate] for t in tiers],
            [L.WHOLE, L.PERCENT])
    for n, period in enumerate(r.phi_rebate, start=1):
        b.heading(f"PHI rebate {period.start:%d/%m/%Y}–{period.end:%d/%m/%Y} (age band {st.age_band}): base, tier 1-3")
        b.table(f"phi{n}", ["Rates", "Base", "Tier 1", "Tier 2", "Tier 3"],
                [["Rebate %", *period.rates[st.age_band]]], ["@", L.RATE, L.RATE, L.RATE, L.RATE])
    b.heading(f"HELP repayment ({r.help.method.replace('_', ' ')})")
    if r.help.method == "marginal":
        b.table("help", ["Repayment income over", "Base", "Rate"], [[x.from_, x.base, x.rate] for x in r.help.bands],
                [L.WHOLE, L.NUMBER, L.PERCENT])
        b.value("help_flat_from", "Flat rate applies from", r.help.flat_from, L.WHOLE)
        b.value("help_flat_rate", "Flat rate on total repayment income", r.help.flat_rate, L.PERCENT)
    else:
        b.table("help", ["Repayment income from", "Rate on total income"], [[x.from_, x.rate] for x in r.help.bands],
                [L.WHOLE, L.PERCENT])
    b.heading("Other rates (for reference)")
    b.value("car_rate", "Car cents per km", r.car.rate)
    b.value("car_cap", "Cents per km business km limit", r.car.km_cap, L.WHOLE)
    b.value("wfh_rate", "Working from home fixed rate per hour", r.wfh_rate)
    b.row += 1
    for name, url in r.sources.items():
        ws.cell(row=b.row, column=1, value=f"Source: {name}")
        ws.cell(row=b.row, column=2, value=url)
        b.row += 1
    L.set_widths(ws, L.RATES_WIDTHS)
    L.setup_print(ws, landscape=False)
    return b.refs


# --- tax calc ------------------------------------------------------------------------------

def _tax_calc(ws: Worksheet, ctx: WorkpaperContext, s: dict[str, str], r: dict[str, str], wp: WorkpaperMap) -> str:
    st = ctx.state
    _title(ws, f"Estimated tax — {ctx.client_code} {fy_label(ctx.income_year)}",
           "Column B is worked out by the formulas below. Column C is the tool's independent calculation B. "
           "Any difference is highlighted and raised as a High review note.")
    ws["A4"], ws["A5"], ws["A6"] = "Days in income year", "Days without private hospital cover", "HELP debt"
    L.style_input(ws.cell(row=4, column=2, value=days_in_fy(ctx.income_year)), L.WHOLE)
    uncovered = 0 if st.hospital_cover_days is None else max(0, days_in_fy(ctx.income_year) - st.hospital_cover_days)
    L.style_input(ws.cell(row=5, column=2, value=uncovered), L.WHOLE)
    if st.hospital_cover_days is None:
        ws.cell(row=5, column=5, value="Assumed covered all year (see review notes)")
    ws.cell(row=6, column=2, value="Yes" if st.help_debt else "No")
    days, days_uncovered = "$B$4", "$B$5"

    head = 8
    for col, h in enumerate(["Component", "Workpaper (formula)", "Python check (calc B)", "Difference",
                             "How it's worked out"], start=1):
        ws.cell(row=head, column=col, value=h)
    L.style_header_row(ws, head, 5)
    row_of = {name: head + 1 + i for i, (name, _) in enumerate(TAXCALC_ROWS)}
    c = {name: f"$B${row}" for name, row in row_of.items()}

    def S(key: str) -> str:
        return s[key]

    rent = f"({S('rent.gross')}-{S('rent.interest')}-{S('rent.capital_works')}-{S('rent.other')})"
    income = "+".join(S(k) for k in ("income.salary", "income.allowances", "income.interest", "income.div_unfranked",
                                     "income.div_franked", "income.franking_credit", "cgt.net_gain", "income.other"))
    ti = c["taxable_income"]
    mlsi = c["mls_income"]
    bf, br, bb = r["brackets.1"], r["brackets.2"], r["brackets.3"]
    mls_from, mls_rate = r["mls.1"], r["mls.2"]
    tier_pos = f"MATCH(MAX(0,{mlsi}-0.001),{mls_from},1)"

    if st.sapto_status != "none":
        sapto = (f"=ROUND(MAX(0,{r['sapto_max']}-{r['sapto_taper']}*MAX(0,{mlsi}-{r['sapto_shade']})),2)",
                 "Max offset less taper × rebate income over the shade-out threshold")
    else:
        sapto = ("=0", "Not eligible")
    if not st.help_debt:
        help_ = ("=0", "No HELP debt")
    elif ctx.rates.help.method == "marginal":
        ri = c["repayment_income"]
        lk = f"{ri}-0.001"
        help_ = (f"=ROUND(IF({ri}<=0,0,IF({ri}>={r['help_flat_from']},{r['help_flat_rate']}*{ri},"
                 f"LOOKUP({lk},{r['help.1']},{r['help.2']})+({ri}-LOOKUP({lk},{r['help.1']},{r['help.1']}))"
                 f"*LOOKUP({lk},{r['help.1']},{r['help.3']}))),2)",
                 "Marginal: base + rate × income over the threshold; flat % of total income at the top")
    else:
        ri = c["repayment_income"]
        help_ = (f"=ROUND(IF({ri}<INDEX({r['help.1']},1),0,LOOKUP({ri},{r['help.1']},{r['help.2']})*{ri}),2)",
                 "Rate for the income band × total repayment income")

    phi_parts = []
    for n in (1, 2):
        phi_parts.append(f"({S(f'phi.rebate_p{n}')}-ROUND({S(f'phi.premiums_p{n}')}*INDEX({r[f'phi{n}.row1']},{tier_pos}),2))")

    formulas = {
        "total_income": (f"={income}+{rent}+{S('income.business_income')}-{S('income.business_expenses')}",
                         "Income labels + net rent + net capital gain + business net (Summary)"),
        "deductions": ("=" + "+".join(S(k) for k in DEDUCTION_KEYS), "D1–D10, D12, D15 (Summary)"),
        "taxable_income": (f"=MAX(0,ROUNDDOWN({c['total_income']}-{c['deductions']}-{S('loss.l1')},0))",
                           "Total income − deductions − prior-year losses; cents dropped"),
        "basic_tax": (f"=ROUND(LOOKUP({ti},{bf},{bb})+({ti}-LOOKUP({ti},{bf},{bf}))*LOOKUP({ti},{bf},{br}),2)",
                      "Base amount for the bracket + marginal rate × income over the threshold (Rates)"),
        "lito": (f"=ROUND(IF({ti}<={r['lito_full_to']},{r['lito_max']},IF({ti}<={r['lito_t1_to']},"
                 f"{r['lito_max']}-{r['lito_t1_rate']}*({ti}-{r['lito_full_to']}),"
                 f"MAX(0,{r['lito_t2_base']}-{r['lito_t2_rate']}*({ti}-{r['lito_t1_to']})))),2)",
                 "$700 reducing by 5c then 1.5c per $1 (Rates)"),
        "sapto": sapto,
        "offsets_used": (f"=MIN({c['basic_tax']},{c['lito']}+{c['sapto']})", "Offsets cannot exceed the tax"),
        "medicare_levy": (f"=ROUND(IF({ti}<={r['ml_lower']},0,MIN({r['ml_rate']}*{ti},{r['ml_shade']}*({ti}-{r['ml_lower']}))),2)",
                          "2%, shaded in above the low-income threshold"),
        "mls_income": (f"={ti}+{S('it.rfb')}+{S('it.resc')}+MAX(0,-{rent})",
                       "Taxable income + fringe benefits + employer super + net rental loss"),
        "mls": (f"=ROUND(LOOKUP(MAX(0,{mlsi}-0.001),{mls_from},{mls_rate})*({ti}+{S('it.rfb')})*{days_uncovered}/{days},2)",
                "Tier rate × (taxable income + fringe benefits) × days without cover"),
        "repayment_income": (f"={mlsi}", "Same income test items as MLS"),
        "help": help_,
        "phi_adjustment": ("=ROUND(" + "+".join(phi_parts) + ",2)",
                           "Rebate received − entitlement at the client's tier, each period"),
        "franking_credits": (f"={S('income.franking_credit')}", "Item 11 U"),
        "payg_withheld": (f"={S('income.salary_tax_withheld')}", "Item 1 tax withheld"),
        "tfn_withheld": (f"={S('income.interest_tfn_withheld')}+{S('income.div_tfn_withheld')}", "Items 10 M and 11 V"),
        "net_result": (f"=ROUND({c['basic_tax']}-{c['offsets_used']}+{c['medicare_levy']}+{c['mls']}+{c['help']}"
                       f"+{c['phi_adjustment']}-{c['franking_credits']}-{c['payg_withheld']}-{c['tfn_withheld']},2)",
                       "Positive = estimated payable; negative = estimated refund"),
    }
    b_values = ctx.calc_b.components()
    for name, label in TAXCALC_ROWS:
        row = row_of[name]
        ws.cell(row=row, column=1, value=label)
        formula, how = formulas[name]
        cell = ws.cell(row=row, column=2, value=formula)
        (L.style_total if name in ("taxable_income", "net_result") else L.style_formula)(cell)
        L.style_input(ws.cell(row=row, column=3, value=b_values[name]))
        L.style_formula(ws.cell(row=row, column=4, value=f"=B{row}-C{row}"))
        ws.cell(row=row, column=5, value=how)
        wp.taxcalc_cells[name] = f"{_q('Tax Calc')}!B{row}"
    L.flag_nonzero(ws, f"D{head + 1}:D{head + len(TAXCALC_ROWS)}")
    note_row = head + len(TAXCALC_ROWS) + 2
    ws.cell(row=note_row, column=1, value="Estimate only. Xero Tax produces the figure that is lodged; use this to "
            "cross-check it. Items flagged in Review Notes are not included in this estimate.")
    L.style_subtitle(ws.cell(row=note_row, column=1))
    L.set_widths(ws, L.TAXCALC_WIDTHS)
    L.setup_print(ws)
    return f"{_q('Tax Calc')}!$B${row_of['net_result']}"


# --- sources, review notes, cover ------------------------------------------------------------

def _sources(ws: Worksheet, ctx: WorkpaperContext) -> None:
    _title(ws, f"Source documents — {ctx.client_code} {fy_label(ctx.income_year)}",
           "Every filed document. Originals are filed unchanged; the tool keeps no copy of their text.")
    headers = ["Doc ID", "Filed name", "Type", "Category", "Document date", "Date found", "Pages", "Parse status",
               "Store", "Redacted on read", "SHA-256 (first 16)"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=FIRST_ROW - 1, column=col, value=h)
    L.style_header_row(ws, FIRST_ROW - 1, len(headers))
    for row, d in enumerate(ctx.documents, start=FIRST_ROW):
        ws.cell(row=row, column=1, value=d.doc_id)
        name = ws.cell(row=row, column=2, value=d.filed_name)
        name.hyperlink = _link(d.filed_location)
        L.style_link(name)
        ws.cell(row=row, column=3, value=d.doc_type)
        ws.cell(row=row, column=4, value=d.category)
        ws.cell(row=row, column=5, value=d.doc_date).number_format = L.DATE
        ws.cell(row=row, column=6, value="Y" if d.date_detected else "N")
        ws.cell(row=row, column=7, value=d.pages)
        ws.cell(row=row, column=8, value=d.parse_status)
        ws.cell(row=row, column=9, value=d.store)
        ws.cell(row=row, column=10, value=", ".join(f"{k}×{v}" for k, v in sorted(d.redaction_counts.items())))
        ws.cell(row=row, column=11, value=d.sha256[:16])
    ws.freeze_panes = f"A{FIRST_ROW}"
    L.set_widths(ws, L.SOURCES_WIDTHS)
    L.setup_print(ws, header_row=FIRST_ROW - 1)


def _citation_text(ctx: WorkpaperContext, citation_id: str | None) -> str:
    if not citation_id:
        return ""
    entry = ctx.citations.get(citation_id)
    if not entry:
        return f"{citation_id} (citation to be confirmed by reviewer)"
    return entry["ref"] + ("" if entry.get("verified") else " (citation to be confirmed by reviewer)")


def _review_notes(ws: Worksheet, ctx: WorkpaperContext) -> None:
    _title(ws, f"Review notes — {ctx.client_code} {fy_label(ctx.income_year)}",
           "Every flag, assumption and client question, ranked by risk then dollar impact. Respond in column J.")
    headers = ["#", "Risk", "Area", "Note", "$ impact", "For", "Action", "Citation", "Refs", "Response"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=FIRST_ROW - 1, column=col, value=h)
    L.style_header_row(ws, FIRST_ROW - 1, len(headers))
    for i, n in enumerate(ctx.notes, start=1):
        row = FIRST_ROW + i - 1
        ws.cell(row=row, column=1, value=i)
        L.style_risk(ws.cell(row=row, column=2, value=n.risk), n.risk)
        area = n.check_id + (f" · {label_name(ctx.labels, n.label_key)}" if n.label_key else "")
        L.style_wrap(ws.cell(row=row, column=3, value=area))
        L.style_wrap(ws.cell(row=row, column=4, value=n.message))
        if n.dollar_impact is not None:
            L.style_formula(ws.cell(row=row, column=5, value=n.dollar_impact))
        ws.cell(row=row, column=6, value=n.audience)
        ws.cell(row=row, column=7, value=n.action)
        L.style_wrap(ws.cell(row=row, column=8, value=_citation_text(ctx, n.citation_id)))
        ws.cell(row=row, column=9, value=", ".join(n.figure_ids + n.doc_ids))
    ws.freeze_panes = f"A{FIRST_ROW}"
    L.set_widths(ws, L.NOTES_WIDTHS)
    L.setup_print(ws, header_row=FIRST_ROW - 1)


def _cover(ws: Worksheet, ctx: WorkpaperContext, net_ref: str) -> None:
    ws.merge_cells("A1:B2")
    ws["A1"] = BANNER
    L.style_banner(ws["A1"])
    start, end = fy_bounds(ctx.income_year)
    risks = Counter(n.risk for n in ctx.notes)
    rows = [
        ("Client code", ctx.client_code),
        ("Income year", f"{fy_label(ctx.income_year)} ({start:%d %B %Y} – {end:%d %B %Y})"),
        ("Prepared by", ctx.prepared_by),
        ("Prepared on", date.today().strftime("%d/%m/%Y")),
        ("Tool version", ctx.tool_version),
        ("Rates file", f"FY{ctx.rates.income_year}: {ctx.rates.status}"),
        ("Estimated payable / (refund)", f"={net_ref}"),
        ("Review notes", f"{risks['High']} High, {risks['Medium']} Medium, {risks['Low']} Low"),
        ("Reviewed by (registered tax agent)", ""),
        ("Review date", ""),
        ("Sign-off", ""),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        L.style_label(ws.cell(row=i, column=1, value=label))
        cell = ws.cell(row=i, column=2, value=value)
        if isinstance(value, str) and value.startswith("="):
            L.style_total(cell)
        if label == "Rates file" and ctx.rates.status != "verified":
            L.style_warning(cell)
    text_row = 4 + len(rows) + 1
    for line in (
        "Prepared by the Tax Return Assistant from documents supplied by the client.",
        "This is a draft. It must be reviewed and signed off by a registered tax agent before anything is lodged.",
        "Nothing has been lodged with the ATO. The return is prepared and lodged in Xero Tax.",
    ):
        L.style_subtitle(ws.cell(row=text_row, column=1, value=line))
        text_row += 1
    L.set_widths(ws, L.COVER_WIDTHS)
    L.setup_print(ws, landscape=False)


# --- entry point ---------------------------------------------------------------------------

def build_workpaper(path: Path, ctx: WorkpaperContext) -> WorkpaperMap:
    wp = WorkpaperMap()
    docs = {d.doc_id: d for d in ctx.documents}
    by_tab: dict[str, list[Figure]] = defaultdict(list)
    for f in ctx.figures:
        by_tab[ctx.labels[f.label_key].tab].append(f)
    tab_order = sorted(by_tab, key=lambda t: min(ctx.labels[f.label_key].order for f in by_tab[t]))
    has_cgt = bool(by_tab.get(CGT_TAB)) or any(e.confirmed for e in ctx.cgt_events)
    if has_cgt and CGT_TAB not in tab_order:
        tab_order.append(CGT_TAB)
        tab_order.sort(key=lambda t: 317 if t == CGT_TAB else min(ctx.labels[f.label_key].order for f in by_tab[t]))

    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    summary = wb.create_sheet("Summary")
    tab_sheets = {tab: wb.create_sheet(tab) for tab in tab_order}
    tax_calc = wb.create_sheet("Tax Calc")
    rates = wb.create_sheet("Rates")
    sources = wb.create_sheet("Sources")
    notes = wb.create_sheet("Review Notes")

    cgt_refs: dict[str, str] = {}
    tab_of: dict[str, str] = {}
    for tab in tab_order:
        figures = by_tab.get(tab, [])
        if tab == CGT_TAB:
            cgt_refs = _cgt_tab(tab_sheets[tab], figures, ctx, docs, wp)
        else:
            _category_tab(tab_sheets[tab], tab, figures, ctx, docs, wp)
        for f in figures:
            tab_of[f.label_key] = tab
    summary_refs = _summary(summary, ctx, tab_of, cgt_refs, wp)
    rate_refs = _rates(rates, ctx)
    net_ref = _tax_calc(tax_calc, ctx, summary_refs, rate_refs, wp)
    _sources(sources, ctx)
    _review_notes(notes, ctx)
    _cover(cover, ctx, net_ref)

    check_workbook(wb, path.name)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    try:
        os.replace(tmp, path)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise FileLockedError(path) from None
    return wp
