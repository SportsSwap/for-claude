"""Every visual choice in the workpaper lives here.

To re-skin the workpaper to the firm's template, change this module only:
build.py decides what goes in which cell, layout.py decides how it looks.
Convention: blue text marks typed-in inputs (figures, rates), black marks formulas.
"""

from __future__ import annotations

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

NUMBER = '#,##0.00;(#,##0.00);"-"'
WHOLE = '#,##0;(#,##0);"-"'
PERCENT = '0.0%'
RATE = '0.000%'
DATE = 'dd/mm/yyyy'

INK = "1F2933"
INPUT_BLUE = "1F4E9A"
MUTED = "6B7280"
HEADER_FILL = PatternFill("solid", fgColor="E5E7EB")
TOTAL_FILL = PatternFill("solid", fgColor="F3F4F6")
BANNER_FILL = PatternFill("solid", fgColor="B91C1C")
WARN_FILL = PatternFill("solid", fgColor="FEE2E2")
RISK_FILLS = {
    "High": PatternFill("solid", fgColor="FECACA"),
    "Medium": PatternFill("solid", fgColor="FEF3C7"),
    "Low": PatternFill("solid", fgColor="E0F2FE"),
}
_THIN = Side(style="thin", color="9CA3AF")
_DOUBLE = Side(style="double", color=INK)


def style_title(cell) -> None:
    cell.font = Font(bold=True, size=14, color=INK)


def style_subtitle(cell) -> None:
    cell.font = Font(italic=True, size=9, color=MUTED)


def style_banner(cell) -> None:
    cell.font = Font(bold=True, size=18, color="FFFFFF")
    cell.fill = BANNER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_header_row(ws: Worksheet, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color=INK)
        cell.fill = HEADER_FILL
        cell.border = Border(bottom=_THIN)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def style_label(cell) -> None:
    cell.font = Font(bold=True, color=INK)


def style_input(cell, fmt: str = NUMBER) -> None:
    cell.font = Font(color=INPUT_BLUE)
    cell.number_format = fmt


def style_formula(cell, fmt: str = NUMBER) -> None:
    cell.font = Font(color=INK)
    cell.number_format = fmt


def style_total(cell, fmt: str = NUMBER) -> None:
    cell.font = Font(bold=True, color=INK)
    cell.number_format = fmt
    cell.fill = TOTAL_FILL
    cell.border = Border(top=_THIN, bottom=_DOUBLE)


def style_warning(cell) -> None:
    cell.font = Font(bold=True, color="B91C1C")
    cell.fill = WARN_FILL


def style_risk(cell, risk: str) -> None:
    cell.fill = RISK_FILLS.get(risk, TOTAL_FILL)
    cell.font = Font(bold=True, color=INK)


def style_link(cell) -> None:
    cell.font = Font(color=INPUT_BLUE, underline="single")


def style_wrap(cell) -> None:
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def set_widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def flag_nonzero(ws: Worksheet, cell_range: str) -> None:
    """Red fill on any cell in the range that is not zero (to the cent)."""
    first = cell_range.split(":")[0]
    ws.conditional_formatting.add(cell_range, FormulaRule(formula=[f"ABS({first})>=0.01"], fill=WARN_FILL))


CATEGORY_WIDTHS = {"A": 10, "B": 12, "C": 42, "D": 16, "E": 44, "F": 7, "G": 9, "H": 11, "I": 14, "J": 15,
                   "K": 4, "L": 30}
SUMMARY_WIDTHS = {"A": 7, "B": 11, "C": 52, "D": 15, "E": 15, "F": 15, "G": 15, "H": 11, "I": 22, "J": 4}
TAXCALC_WIDTHS = {"A": 44, "B": 18, "C": 18, "D": 13, "E": 70}
RATES_WIDTHS = {"A": 34, "B": 14, "C": 14, "D": 14, "E": 14}
SOURCES_WIDTHS = {"A": 8, "B": 46, "C": 11, "D": 13, "E": 12, "F": 10, "G": 7, "H": 14, "I": 8, "J": 22, "K": 16}
NOTES_WIDTHS = {"A": 5, "B": 9, "C": 30, "D": 80, "E": 13, "F": 11, "G": 12, "H": 34, "I": 20, "J": 40}
COVER_WIDTHS = {"A": 30, "B": 50}
