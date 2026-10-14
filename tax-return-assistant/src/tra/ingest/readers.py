"""Turn a file into redacted pages.

Raw text exists only inside :func:`load` and goes straight to the redactor.
Nothing else in the tool ever sees it.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pdfplumber

from ..privacy.redact import RedactedText, redact_pages

SUPPORTED = (".pdf", ".csv", ".xlsx", ".json")


class UnsupportedFile(Exception):
    def __init__(self, suffix: str):
        self.suffix = suffix
        super().__init__(f"Can't read {suffix or 'files without an extension'} files. Supported: {', '.join(SUPPORTED)}")


@dataclass(frozen=True)
class LoadedDocument:
    redacted: RedactedText
    sha256: str
    page_count: int
    image_only_pages: tuple[int, ...]
    suffix: str


def _pdf_pages(path: Path) -> list[str]:
    with pdfplumber.open(path) as pdf:
        return [page.extract_text() or "" for page in pdf.pages]


def _xlsx_pages(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    lines.append(",".join(cells))
        return ["\n".join(lines)]
    finally:
        wb.close()


def load(path: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise UnsupportedFile(suffix)
    sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    if suffix == ".pdf":
        raw = _pdf_pages(path)
    elif suffix == ".xlsx":
        raw = _xlsx_pages(path)
    else:
        raw = [path.read_text(encoding="utf-8-sig", errors="replace")]
    image_only = tuple(i for i, page in enumerate(raw, start=1) if not page.strip()) if suffix == ".pdf" else ()
    redacted = redact_pages(raw)
    del raw
    return LoadedDocument(redacted=redacted, sha256=sha256, page_count=len(redacted.pages),
                          image_only_pages=image_only, suffix=suffix)
